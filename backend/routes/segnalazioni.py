"""Route: Segnalazioni — i Segnalatori inseriscono clienti segnalati, il BackOffice Commessa li esita."""
import io
import logging
import re
import uuid
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import List, Optional
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from fastapi.responses import Response
from pydantic import BaseModel, Field

from database import db
from security import get_current_user
from models import User, UserRole
from services import NextcloudClient

router = APIRouter()
logger = logging.getLogger(__name__)

ESITI = ["OK", "KO", "NR"]
SOTTO_ESITI_OK = ["OK Energy", "OK Gas", "OK Dual", "OK Fisso", "OK Fisso + Mobile", "OK Mobile"]
ESITO_ROLES = {UserRole.ADMIN, UserRole.RESPONSABILE_COMMESSA, UserRole.BACKOFFICE_COMMESSA}
MAX_DOC_SIZE = 25 * 1024 * 1024
ROME = ZoneInfo("Europe/Rome")


class SegnalazioneCreate(BaseModel):
    nome: str = Field(min_length=1)
    cognome: str = Field(min_length=1)
    telefono: str = Field(min_length=5, max_length=20, pattern=r"^\+?[0-9 ./-]{5,20}$")
    note: Optional[str] = ""


class AssegnazioneUpdate(BaseModel):
    commessa_id: str
    servizio_id: str


class EsitoUpdate(BaseModel):
    esito: str
    sotto_esiti: List[str] = []
    nota_esito: Optional[str] = ""


def _scope_query(user: User) -> dict:
    if user.role == UserRole.ADMIN:
        return {}
    if user.role == UserRole.SEGNALATORE:
        return {"created_by": user.id}
    if user.role in ESITO_ROLES:
        q = {"commessa_id": {"$in": user.commesse_autorizzate or []}}
        if user.servizi_autorizzati:
            q["servizio_id"] = {"$in": user.servizi_autorizzati}
        return {"$or": [{"commessa_id": None}, q]}
    raise HTTPException(status_code=403, detail="Accesso non consentito alle segnalazioni")


async def _get_accessible(segnalazione_id: str, user: User) -> dict:
    doc = await db.segnalazioni.find_one({"id": segnalazione_id, **_scope_query(user)}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Segnalazione non trovata")
    return doc


async def _aruba_commessa_for(user: User) -> Optional[str]:
    """Prima commessa del segnalatore con Aruba Drive attivo (cartella dove salvare i documenti)."""
    ids = user.commesse_autorizzate or []
    found = await db.commesse.find({"id": {"$in": ids}, "aruba_drive_config.enabled": True}, {"_id": 0, "id": 1}).to_list(None)
    enabled = {c["id"] for c in found}
    return next((i for i in ids if i in enabled), ids[0] if ids else None)


async def _aruba_client(commessa_id: Optional[str]) -> NextcloudClient:
    commessa = await db.commesse.find_one({"id": commessa_id}, {"_id": 0}) if commessa_id else None
    cfg = (commessa or {}).get("aruba_drive_config") or {}
    if not cfg.get("enabled"):
        settings = await db.settings.find_one({"key": "aruba_drive_global"}, {"_id": 0})
        cfg = (settings or {}).get("value") or {}
    if not cfg.get("enabled") or not cfg.get("url"):
        raise HTTPException(status_code=400, detail="Aruba Drive non è configurato per questa commessa. Contattare l'amministratore.")
    root = (cfg.get("root_folder_path") or (commessa or {}).get("nome") or "Documenti").strip("/")
    return NextcloudClient(base_url=cfg["url"], username=cfg.get("username"), password=cfg.get("password"),
                           folder_path=f"{root}/Segnalazioni")


def _safe(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", (s or "").strip()).strip("_")


@router.get("/segnalazioni/options")
async def get_segnalazioni_options(current_user: User = Depends(get_current_user)):
    """Commesse e servizi su cui l'utente può inserire/filtrare segnalazioni."""
    _scope_query(current_user)
    c_query = {"is_active": {"$ne": False}}
    s_query = {"is_active": {"$ne": False}}
    if current_user.role != UserRole.ADMIN:
        c_query["id"] = {"$in": current_user.commesse_autorizzate or []}
        if current_user.servizi_autorizzati:
            s_query["id"] = {"$in": current_user.servizi_autorizzati}
    commesse = await db.commesse.find(c_query, {"_id": 0, "id": 1, "nome": 1}).to_list(None)
    s_query["commessa_id"] = {"$in": [c["id"] for c in commesse]}
    servizi = await db.servizi.find(s_query, {"_id": 0, "id": 1, "nome": 1, "commessa_id": 1}).to_list(None)
    result = []
    for c in sorted(commesse, key=lambda x: (x.get("nome") or "").lower()):
        srv = sorted([s for s in servizi if s["commessa_id"] == c["id"]], key=lambda x: (x.get("nome") or "").lower())
        if srv:
            result.append({"id": c["id"], "nome": c["nome"], "servizi": [{"id": s["id"], "nome": s["nome"]} for s in srv]})
    return {"commesse": result, "esiti": ESITI, "sotto_esiti_ok": SOTTO_ESITI_OK}


@router.get("/segnalazioni/da-lavorare/count")
async def count_da_lavorare(current_user: User = Depends(get_current_user)):
    if current_user.role not in ESITO_ROLES:
        return {"count": 0}
    return {"count": await db.segnalazioni.count_documents({**_scope_query(current_user), "esito": None})}


@router.post("/segnalazioni")
async def create_segnalazione(data: SegnalazioneCreate, current_user: User = Depends(get_current_user)):
    if current_user.role not in (UserRole.SEGNALATORE, UserRole.ADMIN):
        raise HTTPException(status_code=403, detail="Solo i Segnalatori possono inserire segnalazioni")
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "id": str(uuid.uuid4()),
        "nome": data.nome.strip(),
        "cognome": data.cognome.strip(),
        "telefono": data.telefono.strip(),
        "note": (data.note or "").strip(),
        "commessa_id": None,
        "commessa_nome": None,
        "servizio_id": None,
        "servizio_nome": None,
        "aruba_commessa_id": await _aruba_commessa_for(current_user),
        "esito": None,
        "in_carico_id": None,
        "in_carico_username": None,
        "sotto_esiti": [],
        "nota_esito": "",
        "esito_history": [],
        "documents": [],
        "created_by": current_user.id,
        "created_by_username": current_user.username,
        "created_at": now,
        "updated_at": now,
    }
    await db.segnalazioni.insert_one(doc)
    doc.pop("_id", None)
    return doc


def _day_to_utc_iso(day: str, end: bool = False) -> str:
    try:
        d = datetime.strptime(day, "%Y-%m-%d").replace(tzinfo=ROME)
    except ValueError:
        raise HTTPException(status_code=400, detail="Formato data non valido (YYYY-MM-DD)")
    if end:
        d += timedelta(days=1)
    return d.astimezone(timezone.utc).isoformat()


def _filters_query(
    current_user: User = Depends(get_current_user),
    esito: Optional[str] = None,
    commessa_id: Optional[str] = None,
    servizio_id: Optional[str] = None,
    segnalatore_id: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    search: Optional[str] = None,
) -> dict:
    f = {}
    if esito == "da_lavorare":
        f["esito"] = None
    elif esito in ESITI:
        f["esito"] = esito
    if commessa_id:
        f["commessa_id"] = None if commessa_id == "none" else commessa_id
    if servizio_id:
        f["servizio_id"] = servizio_id
    if segnalatore_id:
        f["created_by"] = segnalatore_id
    if date_from or date_to:
        f["created_at"] = {}
        if date_from:
            f["created_at"]["$gte"] = _day_to_utc_iso(date_from)
        if date_to:
            f["created_at"]["$lt"] = _day_to_utc_iso(date_to, end=True)
    if search and search.strip():
        rx = {"$regex": re.escape(search.strip()), "$options": "i"}
        f["$or"] = [{"nome": rx}, {"cognome": rx}, {"telefono": rx}, {"created_by_username": rx}]
    return {"$and": [_scope_query(current_user), f]}


@router.get("/segnalazioni")
async def list_segnalazioni(query: dict = Depends(_filters_query)):
    return await db.segnalazioni.find(query, {"_id": 0, "esito_history": 0}).sort("created_at", -1).to_list(1000)


@router.get("/segnalazioni/stats/summary")
async def segnalazioni_stats(query: dict = Depends(_filters_query)):
    docs = await db.segnalazioni.find(
        query, {"_id": 0, "id": 1, "nome": 1, "cognome": 1, "esito": 1, "sotto_esiti": 1, "esito_at": 1, "servizio_nome": 1}
    ).to_list(None)
    by_esito = {e: sum(1 for d in docs if d.get("esito") == e) for e in ESITI}
    esitate = sum(by_esito.values())
    sotto = {s: sum(1 for d in docs if s in (d.get("sotto_esiti") or [])) for s in SOTTO_ESITI_OK}
    recenti = sorted([d for d in docs if d.get("esito_at")], key=lambda d: d["esito_at"], reverse=True)[:5]
    return {
        "totale": len(docs),
        "da_lavorare": len(docs) - esitate,
        "esitate": esitate,
        "per_esito": by_esito,
        "tasso_ok": round(by_esito["OK"] * 100 / esitate, 1) if esitate else 0,
        "sotto_esiti": sotto,
        "ultime_esitate": recenti,
    }


@router.get("/segnalazioni/segnalatori/list")
async def list_segnalatori(current_user: User = Depends(get_current_user)):
    if current_user.role == UserRole.SEGNALATORE:
        return []
    if current_user.role == UserRole.ADMIN:
        users_q = {"role": UserRole.SEGNALATORE.value}
    else:
        users_q = {"id": {"$in": await db.segnalazioni.distinct("created_by", _scope_query(current_user))}}
    users = await db.users.find(users_q, {"_id": 0, "id": 1, "username": 1}).to_list(None)
    return sorted(users, key=lambda u: u["username"].lower())


def _rome(iso: Optional[str]) -> str:
    return datetime.fromisoformat(iso).astimezone(ROME).strftime("%d/%m/%Y %H:%M") if iso else ""


def _cell(v):
    return f"'{v}" if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


@router.get("/segnalazioni/export/excel")
async def export_segnalazioni_excel(
    query: dict = Depends(_filters_query), current_user: User = Depends(get_current_user)
):
    if current_user.role != UserRole.ADMIN:
        raise HTTPException(status_code=403, detail="Export disponibile solo per Admin")
    docs = await db.segnalazioni.find(query, {"_id": 0, "esito_history": 0}).sort("created_at", -1).to_list(None)
    wb = Workbook()
    ws = wb.active
    ws.title = "Segnalazioni"
    headers = ["Data segnalazione", "Nome", "Cognome", "Telefono", "Note", "Commessa", "Servizio", "Segnalatore",
               "Esito", "Sotto esiti", "Nota esito", "Esitato da", "Data esito", "N. documenti"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
    for d in docs:
        ws.append([_cell(v) for v in [_rome(d.get("created_at")), d.get("nome"), d.get("cognome"), d.get("telefono"), d.get("note"),
                   d.get("commessa_nome"), d.get("servizio_nome"), d.get("created_by_username"),
                   d.get("esito") or "Da lavorare", ", ".join(d.get("sotto_esiti") or []), d.get("nota_esito"),
                   d.get("esito_by") or "", _rome(d.get("esito_at")), len(d.get("documents") or [])]])
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 4)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    name = f"segnalazioni_{datetime.now(ROME).strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/segnalazioni/{segnalazione_id}")
async def get_segnalazione(segnalazione_id: str, current_user: User = Depends(get_current_user)):
    return await _get_accessible(segnalazione_id, current_user)


def _check_in_carico(seg: dict, user: User):
    if user.role == UserRole.ADMIN:
        return
    if not seg.get("in_carico_id"):
        raise HTTPException(status_code=409, detail="Prendi in carico la segnalazione prima di lavorarla")
    if seg["in_carico_id"] != user.id:
        raise HTTPException(status_code=409, detail=f"Segnalazione già in carico a {seg.get('in_carico_username')}")


def _lock_filter(segnalazione_id: str, user: User) -> dict:
    return {"id": segnalazione_id} if user.role == UserRole.ADMIN else {"id": segnalazione_id, "in_carico_id": user.id}


@router.post("/segnalazioni/{segnalazione_id}/presa-in-carico")
async def prendi_in_carico(segnalazione_id: str, current_user: User = Depends(get_current_user)):
    if current_user.role not in ESITO_ROLES:
        raise HTTPException(status_code=403, detail="Solo il BackOffice Commessa può prendere in carico")
    await _get_accessible(segnalazione_id, current_user)
    now = datetime.now(timezone.utc).isoformat()
    cond = {"id": segnalazione_id}
    if current_user.role != UserRole.ADMIN:
        cond["$or"] = [{"in_carico_id": None}, {"in_carico_id": {"$exists": False}}, {"in_carico_id": current_user.id}]
    res = await db.segnalazioni.update_one(cond, {"$set": {
        "in_carico_id": current_user.id, "in_carico_username": current_user.username, "in_carico_at": now, "updated_at": now}})
    if res.matched_count == 0:
        seg = await db.segnalazioni.find_one({"id": segnalazione_id}, {"_id": 0, "in_carico_username": 1})
        raise HTTPException(status_code=409, detail=f"Segnalazione già in carico a {(seg or {}).get('in_carico_username')}")
    return await db.segnalazioni.find_one({"id": segnalazione_id}, {"_id": 0})


@router.put("/segnalazioni/{segnalazione_id}/assegnazione")
async def update_assegnazione(segnalazione_id: str, data: AssegnazioneUpdate, current_user: User = Depends(get_current_user)):
    if current_user.role not in ESITO_ROLES:
        raise HTTPException(status_code=403, detail="Solo il BackOffice Commessa può assegnare commessa e servizio")
    seg = await _get_accessible(segnalazione_id, current_user)
    _check_in_carico(seg, current_user)
    servizio = await db.servizi.find_one({"id": data.servizio_id, "commessa_id": data.commessa_id, "is_active": {"$ne": False}}, {"_id": 0, "nome": 1})
    if not servizio:
        raise HTTPException(status_code=400, detail="Servizio non valido per la commessa selezionata")
    if current_user.role != UserRole.ADMIN and (
        data.commessa_id not in (current_user.commesse_autorizzate or [])
        or (current_user.servizi_autorizzati and data.servizio_id not in current_user.servizi_autorizzati)
    ):
        raise HTTPException(status_code=403, detail="Commessa o servizio non autorizzati")
    commessa = await db.commesse.find_one({"id": data.commessa_id}, {"_id": 0, "nome": 1})
    update = {"commessa_id": data.commessa_id, "commessa_nome": (commessa or {}).get("nome"),
              "servizio_id": data.servizio_id, "servizio_nome": servizio.get("nome"),
              "assegnato_by": current_user.username, "updated_at": datetime.now(timezone.utc).isoformat()}
    if not seg.get("aruba_commessa_id"):
        update["aruba_commessa_id"] = data.commessa_id
    res = await db.segnalazioni.update_one(_lock_filter(segnalazione_id, current_user), {"$set": update, "$push": {"assegnazione_history": {
        "commessa_nome": update["commessa_nome"], "servizio_nome": update["servizio_nome"],
        "by_username": current_user.username, "at": update["updated_at"]}}})
    if res.matched_count == 0:
        raise HTTPException(status_code=409, detail="La segnalazione è stata presa in carico da un altro utente")
    return await db.segnalazioni.find_one({"id": segnalazione_id}, {"_id": 0})


@router.put("/segnalazioni/{segnalazione_id}/esito")
async def update_esito(segnalazione_id: str, data: EsitoUpdate, current_user: User = Depends(get_current_user)):
    if current_user.role not in ESITO_ROLES:
        raise HTTPException(status_code=403, detail="Solo il BackOffice Commessa può esitare le segnalazioni")
    seg = await _get_accessible(segnalazione_id, current_user)
    _check_in_carico(seg, current_user)
    if not seg.get("commessa_id") or not seg.get("servizio_id"):
        raise HTTPException(status_code=400, detail="Assegna commessa e servizio prima di dare l'esito")
    if data.esito not in ESITI:
        raise HTTPException(status_code=400, detail="Esito non valido")
    sotto = [s for s in SOTTO_ESITI_OK if s in data.sotto_esiti] if data.esito == "OK" else []
    if data.esito == "OK" and not sotto:
        raise HTTPException(status_code=400, detail="Con esito OK seleziona almeno un sotto esito")
    now = datetime.now(timezone.utc).isoformat()
    entry = {"esito": data.esito, "sotto_esiti": sotto, "nota_esito": (data.nota_esito or "").strip(),
             "by_id": current_user.id, "by_username": current_user.username, "at": now}
    res = await db.segnalazioni.update_one(
        _lock_filter(segnalazione_id, current_user),
        {"$set": {"esito": data.esito, "sotto_esiti": sotto, "nota_esito": entry["nota_esito"],
                  "esito_by": current_user.username, "esito_at": now, "updated_at": now,
                  "in_carico_id": None, "in_carico_username": None, "in_carico_at": None},
         "$push": {"esito_history": entry}},
    )
    if res.matched_count == 0:
        raise HTTPException(status_code=409, detail="La segnalazione è stata presa in carico da un altro utente")
    return await db.segnalazioni.find_one({"id": segnalazione_id}, {"_id": 0})


@router.post("/segnalazioni/{segnalazione_id}/documents")
async def upload_segnalazione_document(
    segnalazione_id: str, file: UploadFile = File(...), current_user: User = Depends(get_current_user)
):
    seg = await _get_accessible(segnalazione_id, current_user)
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="File vuoto")
    if len(content) > MAX_DOC_SIZE:
        raise HTTPException(status_code=400, detail="File troppo grande (max 25MB)")
    aruba_commessa_id = seg.get("aruba_commessa_id") or seg.get("commessa_id")
    client = await _aruba_client(aruba_commessa_id)
    original = file.filename or "documento"
    stem, ext = _safe(Path(original).stem) or "documento", Path(original).suffix.lower()
    prefix = "_".join(p for p in [_safe(seg["telefono"]), _safe(seg["nome"]), _safe(seg["cognome"])] if p)
    filename = f"{prefix}_{uuid.uuid4().hex[:6]}_{stem}{ext}"
    ok, cloud_path = await client.upload_file(content, filename)
    if not ok:
        raise HTTPException(status_code=503, detail="Errore di caricamento su Aruba Drive. Il documento NON è stato salvato.")
    doc = {
        "id": str(uuid.uuid4()),
        "filename": filename,
        "original_filename": original,
        "cloud_path": cloud_path,
        "file_type": file.content_type or "application/octet-stream",
        "file_size": len(content),
        "aruba_commessa_id": aruba_commessa_id,
        "uploaded_by": current_user.username,
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.segnalazioni.update_one({"id": segnalazione_id}, {"$push": {"documents": doc}})
    return doc


@router.get("/segnalazioni/{segnalazione_id}/documents/{document_id}/download")
async def download_segnalazione_document(
    segnalazione_id: str, document_id: str, current_user: User = Depends(get_current_user)
):
    seg = await _get_accessible(segnalazione_id, current_user)
    doc = next((d for d in seg.get("documents", []) if d["id"] == document_id), None)
    if not doc:
        raise HTTPException(status_code=404, detail="Documento non trovato")
    client = await _aruba_client(doc.get("aruba_commessa_id") or seg.get("aruba_commessa_id") or seg.get("commessa_id"))
    ok, content = await client.download_file(doc["filename"])
    if not ok:
        raise HTTPException(status_code=404, detail="File non trovato su Aruba Drive")
    return Response(content=content, media_type=doc["file_type"],
                    headers={"Content-Disposition": f'attachment; filename="{doc["original_filename"]}"'})
