# Segnalatori iteration 29: Excel export (admin only), segnalatore filter, date range filter, stats summary
import io
import os
from datetime import datetime, timezone, timedelta
from zoneinfo import ZoneInfo

import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
FASTWEB_ID = "4cb70f28-6278-4d0f-b2b7-65f2b783f3f1"
TLS_ID = "cc0648c1-0df1-4530-8281-f4c940934916"
ROME = ZoneInfo("Europe/Rome")
HEADERS = ["Data segnalazione", "Nome", "Cognome", "Telefono", "Note", "Commessa", "Servizio", "Segnalatore",
           "Esito", "Sotto esiti", "Competenze", "Nota esito", "Esitato da", "Data esito", "N. documenti"]


def _take(sid, headers):
    """BO must take the segnalazione in carico before working it (presa in carico lock)."""
    r = requests.post(f"{API}/segnalazioni/{sid}/presa-in-carico", headers=headers)
    assert r.status_code == 200, r.text


def _login(u, p):
    r = requests.post(f"{API}/auth/login", json={"username": u, "password": p})
    if r.status_code != 200:
        pytest.skip(f"login failed {u}: {r.status_code}")
    return {"Authorization": f"Bearer {r.json()['access_token']}"}, r.json().get("user", {})


@pytest.fixture(scope="module")
def admin():
    return _login("admin", "admin123")[0]


@pytest.fixture(scope="module")
def segn_full():
    return _login("segnalatore_test", "test12345")


@pytest.fixture(scope="module")
def segn(segn_full):
    return segn_full[0]


@pytest.fixture(scope="module")
def bo():
    return _login("bo_segn_test", "test12345")[0]


@pytest.fixture(scope="module")
def segn_id(admin):
    r = requests.get(f"{API}/segnalazioni/segnalatori/list", headers=admin)
    u = next((x for x in r.json() if x["username"] == "segnalatore_test"), None)
    assert u, r.text
    return u["id"]


@pytest.fixture(scope="module")
def created(segn, admin):
    r = requests.post(f"{API}/segnalazioni", headers=segn, json={
        "nome": "TEST_Export", "cognome": "TEST_Stats", "telefono": "3339998887",
        "note": "pytest export"})
    assert r.status_code == 200, r.text
    a = requests.put(f"{API}/segnalazioni/{r.json()['id']}/assegnazione", headers=admin,
                     json={"commessa_id": FASTWEB_ID, "servizio_id": TLS_ID})
    assert a.status_code == 200, a.text
    yield a.json()
    try:
        from pymongo import MongoClient
        from dotenv import dotenv_values
        env = dotenv_values("/app/backend/.env")
        MongoClient(env["MONGO_URL"])[env["DB_NAME"]].segnalazioni.delete_many({"nome": {"$regex": "^TEST_"}})
    except Exception as e:
        print("cleanup failed", e)


def _today_rome():
    return datetime.now(ROME).strftime("%Y-%m-%d")


class TestExport:
    def test_admin_export_xlsx(self, admin, created):
        r = requests.get(f"{API}/segnalazioni/export/excel", headers=admin)
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers["content-type"]
        assert ".xlsx" in r.headers.get("content-disposition", "")
        ws = load_workbook(io.BytesIO(r.content)).active
        rows = list(ws.iter_rows(values_only=True))
        assert list(rows[0]) == HEADERS
        names = [row[1] for row in rows[1:]]
        assert "TEST_Export" in names
        # count matches list
        lst = requests.get(f"{API}/segnalazioni", headers=admin).json()
        assert len(rows) - 1 == len(lst)

    def test_export_respects_filters(self, admin, created):
        r = requests.get(f"{API}/segnalazioni/export/excel", headers=admin, params={"esito": "da_lavorare", "search": "TEST_Export"})
        rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
        assert len(rows) == 2
        assert rows[1][8] == "Da lavorare" and rows[1][7] == "segnalatore_test"
        r2 = requests.get(f"{API}/segnalazioni/export/excel", headers=admin, params={"date_from": "2020-01-01", "date_to": "2020-01-02"})
        assert len(list(load_workbook(io.BytesIO(r2.content)).active.iter_rows(values_only=True))) == 1

    def test_export_forbidden_segn(self, segn):
        r = requests.get(f"{API}/segnalazioni/export/excel", headers=segn)
        assert r.status_code == 403

    def test_export_forbidden_bo(self, bo):
        r = requests.get(f"{API}/segnalazioni/export/excel", headers=bo)
        assert r.status_code == 403

    def test_export_unauth(self):
        r = requests.get(f"{API}/segnalazioni/export/excel")
        assert r.status_code in (401, 403)


class TestSegnalatoreFilter:
    def test_list_admin(self, admin, segn_id):
        r = requests.get(f"{API}/segnalazioni/segnalatori/list", headers=admin)
        assert r.status_code == 200 and any(u["id"] == segn_id for u in r.json())
        assert all("_id" not in u for u in r.json())

    def test_list_bo(self, bo, segn_id):
        r = requests.get(f"{API}/segnalazioni/segnalatori/list", headers=bo)
        assert r.status_code == 200 and any(u["id"] == segn_id for u in r.json())

    def test_list_segn_empty(self, segn):
        r = requests.get(f"{API}/segnalazioni/segnalatori/list", headers=segn)
        assert r.status_code == 200 and r.json() == []

    def test_filter_by_segnalatore(self, admin, segn_id, created):
        r = requests.get(f"{API}/segnalazioni", headers=admin, params={"segnalatore_id": segn_id})
        assert r.status_code == 200 and len(r.json()) >= 1
        assert all(s["created_by"] == segn_id for s in r.json())
        r2 = requests.get(f"{API}/segnalazioni", headers=admin, params={"segnalatore_id": "nonexistent"})
        assert r2.json() == []

    def test_segn_cannot_see_others_via_filter(self, segn, admin):
        # segnalatore passing another id must get nothing (scope AND)
        r = requests.get(f"{API}/segnalazioni", headers=segn, params={"segnalatore_id": "someone-else"})
        assert r.json() == []


class TestDateFilter:
    def test_today_includes_created(self, admin, created):
        t = _today_rome()
        r = requests.get(f"{API}/segnalazioni", headers=admin, params={"date_from": t, "date_to": t})
        assert r.status_code == 200
        assert any(s["id"] == created["id"] for s in r.json())

    def test_past_range_excludes(self, admin, created):
        r = requests.get(f"{API}/segnalazioni", headers=admin, params={"date_from": "2020-01-01", "date_to": "2020-12-31"})
        assert r.status_code == 200 and all(s["id"] != created["id"] for s in r.json())

    def test_only_from_tomorrow_excludes(self, admin, created):
        tomorrow = (datetime.now(ROME) + timedelta(days=1)).strftime("%Y-%m-%d")
        r = requests.get(f"{API}/segnalazioni", headers=admin, params={"date_from": tomorrow})
        assert all(s["id"] != created["id"] for s in r.json())

    def test_only_to_today_includes(self, admin, created):
        r = requests.get(f"{API}/segnalazioni", headers=admin, params={"date_to": _today_rome()})
        assert any(s["id"] == created["id"] for s in r.json())

    def test_invalid_date_400(self, admin):
        for p in ({"date_from": "27/10/2026"}, {"date_to": "abc"}):
            r = requests.get(f"{API}/segnalazioni", headers=admin, params=p)
            assert r.status_code == 400, (p, r.status_code)
            r = requests.get(f"{API}/segnalazioni/stats/summary", headers=admin, params=p)
            assert r.status_code == 400


class TestStats:
    def test_stats_admin_consistent(self, admin, created):
        st = requests.get(f"{API}/segnalazioni/stats/summary", headers=admin).json()
        lst = requests.get(f"{API}/segnalazioni", headers=admin).json()
        assert st["totale"] == len(lst)
        assert st["da_lavorare"] == sum(1 for s in lst if not s.get("esito"))
        for e in ("OK", "KO", "NR"):
            assert st["per_esito"][e] == sum(1 for s in lst if s.get("esito") == e)
        assert st["esitate"] == st["totale"] - st["da_lavorare"]
        assert set(st["sotto_esiti"].keys()) == {"OK Energy", "OK Gas", "OK Dual", "OK Fisso", "OK Fisso + Mobile", "OK Mobile"}
        assert len(st["ultime_esitate"]) <= 5

    def test_stats_segn_own_only(self, segn, segn_full, created):
        st = requests.get(f"{API}/segnalazioni/stats/summary", headers=segn).json()
        lst = requests.get(f"{API}/segnalazioni", headers=segn).json()
        assert st["totale"] == len(lst)
        assert all(s["created_by_username"] == "segnalatore_test" for s in lst)

    def test_stats_follow_date(self, admin, created):
        st = requests.get(f"{API}/segnalazioni/stats/summary", headers=admin, params={"date_from": "2020-01-01", "date_to": "2020-01-02"}).json()
        assert st["totale"] == 0 and st["tasso_ok"] == 0 and st["ultime_esitate"] == []
        t = _today_rome()
        st2 = requests.get(f"{API}/segnalazioni/stats/summary", headers=admin, params={"date_from": t, "date_to": t}).json()
        assert st2["totale"] >= 1 and st2["da_lavorare"] >= 1

    def test_stats_after_esito(self, bo, segn, created):
        before = requests.get(f"{API}/segnalazioni/stats/summary", headers=segn).json()
        _take(created['id'], bo)
        r = requests.put(f"{API}/segnalazioni/{created['id']}/esito", headers=bo,
                         json={"esito": "OK", "sotto_esiti": ["OK Energy"], "nota_esito": "pytest"})
        assert r.status_code == 200, r.text
        after = requests.get(f"{API}/segnalazioni/stats/summary", headers=segn).json()
        assert after["per_esito"]["OK"] == before["per_esito"]["OK"] + 1
        assert after["sotto_esiti"]["OK Energy"] == before["sotto_esiti"]["OK Energy"] + 1
        assert after["da_lavorare"] == before["da_lavorare"] - 1
        assert after["ultime_esitate"][0]["id"] == created["id"]
        exp = round(after["per_esito"]["OK"] * 100 / after["esitate"], 1)
        assert after["tasso_ok"] == exp
