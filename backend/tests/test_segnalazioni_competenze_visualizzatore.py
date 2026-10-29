# Segnalatori: Competenze (admin only, on OK sotto esiti) + Visualizzatore role (read-only full view)
import io
import os
import uuid
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
FASTWEB_ID = "4cb70f28-6278-4d0f-b2b7-65f2b783f3f1"
TLS_ID = "cc0648c1-0df1-4530-8281-f4c940934916"
ASSIGN = {"commessa_id": FASTWEB_ID, "servizio_id": TLS_ID}
CREATED_USERS = []


def _login(u, p):
    r = requests.post(f"{API}/auth/login", json={"username": u, "password": p})
    if r.status_code != 200:
        pytest.skip(f"login failed {u}: {r.status_code}")
    j = r.json()
    return {"Authorization": f"Bearer {j['access_token']}"}


@pytest.fixture(scope="module")
def admin():
    return _login("admin", "admin123")


@pytest.fixture(scope="module")
def segn():
    return _login("segnalatore_test", "test12345")


@pytest.fixture(scope="module")
def bo():
    return _login("bo_segn_test", "test12345")


@pytest.fixture(scope="module")
def vis():
    return _login("visualizzatore_test", "test12345")


@pytest.fixture(scope="module", autouse=True)
def cleanup():
    yield
    from pymongo import MongoClient
    from dotenv import dotenv_values
    env = dotenv_values("/app/backend/.env")
    db = MongoClient(env["MONGO_URL"])[env["DB_NAME"]]
    db.segnalazioni.delete_many({"cognome": "TEST_COMP"})
    db.users.delete_many({"username": {"$regex": "^test_vis_"}})


def _new(segn, name):
    r = requests.post(f"{API}/segnalazioni", headers=segn, json={"nome": f"TEST_{name}", "cognome": "TEST_COMP", "telefono": "3339998888"})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _esito(bo, sid, esito, sotto=None, assign=True):
    assert requests.post(f"{API}/segnalazioni/{sid}/presa-in-carico", headers=bo).status_code == 200
    if assign:
        assert requests.put(f"{API}/segnalazioni/{sid}/assegnazione", headers=bo, json=ASSIGN).status_code == 200
    r = requests.put(f"{API}/segnalazioni/{sid}/esito", headers=bo, json={"esito": esito, "sotto_esiti": sotto or []})
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture(scope="module")
def ok_seg(segn, bo):
    sid = _new(segn, "OkComp")
    _esito(bo, sid, "OK", ["OK Gas", "OK Energy"])
    return sid


class TestCompetenze:
    def test_options_include_competenze(self, admin):
        r = requests.get(f"{API}/segnalazioni/options", headers=admin)
        assert r.status_code == 200 and r.json()["competenze"] == ["Rea", "Andrea"]

    def test_admin_set_and_persist(self, admin, ok_seg):
        r = requests.put(f"{API}/segnalazioni/{ok_seg}/competenze", headers=admin, json={"competenze": {"OK Gas": "Rea", "OK Energy": "Andrea"}})
        assert r.status_code == 200, r.text
        d = requests.get(f"{API}/segnalazioni/{ok_seg}", headers=admin).json()
        assert d["competenze"] == {"OK Gas": "Rea", "OK Energy": "Andrea"}
        assert d["competenze_by"] == "admin"

    def test_reset_to_none_removes(self, admin, ok_seg):
        r = requests.put(f"{API}/segnalazioni/{ok_seg}/competenze", headers=admin, json={"competenze": {"OK Gas": "Rea", "OK Energy": ""}})
        assert r.status_code == 200
        d = requests.get(f"{API}/segnalazioni/{ok_seg}", headers=admin).json()
        assert d["competenze"] == {"OK Gas": "Rea"}

    def test_invalid_value_400(self, admin, ok_seg):
        r = requests.put(f"{API}/segnalazioni/{ok_seg}/competenze", headers=admin, json={"competenze": {"OK Gas": "Mario"}})
        assert r.status_code == 400

    def test_invalid_key_400(self, admin, ok_seg):
        r = requests.put(f"{API}/segnalazioni/{ok_seg}/competenze", headers=admin, json={"competenze": {"OK Mobile": "Rea"}})
        assert r.status_code == 400

    def test_non_ok_400(self, admin, segn, bo):
        sid = _new(segn, "KoComp")
        _esito(bo, sid, "KO")
        r = requests.put(f"{API}/segnalazioni/{sid}/competenze", headers=admin, json={"competenze": {}})
        assert r.status_code == 400
        sid2 = _new(segn, "NoEsito")
        assert requests.put(f"{API}/segnalazioni/{sid2}/competenze", headers=admin, json={"competenze": {}}).status_code == 400

    def test_bo_vis_segn_403(self, bo, vis, segn, ok_seg):
        for h in (bo, vis, segn):
            r = requests.put(f"{API}/segnalazioni/{ok_seg}/competenze", headers=h, json={"competenze": {"OK Gas": "Andrea"}})
            assert r.status_code == 403

    def test_bo_sees_competenze(self, bo, ok_seg):
        d = requests.get(f"{API}/segnalazioni/{ok_seg}", headers=bo).json()
        assert d["competenze"] == {"OK Gas": "Rea"}
        lst = requests.get(f"{API}/segnalazioni", headers=bo).json()
        assert next(s for s in lst if s["id"] == ok_seg)["competenze"] == {"OK Gas": "Rea"}

    def test_segnalatore_never_sees(self, segn, ok_seg):
        d = requests.get(f"{API}/segnalazioni/{ok_seg}", headers=segn).json()
        assert "competenze" not in d and "in_carico_id" not in d
        lst = requests.get(f"{API}/segnalazioni", headers=segn).json()
        row = next(s for s in lst if s["id"] == ok_seg)
        assert "competenze" not in row
        if "competenze_by" in d:
            pytest.fail("competenze_by leaked to segnalatore")

    def test_esito_change_drops_competenza(self, admin, bo, ok_seg):
        requests.put(f"{API}/segnalazioni/{ok_seg}/competenze", headers=admin, json={"competenze": {"OK Gas": "Rea", "OK Energy": "Andrea"}})
        _esito(bo, ok_seg, "OK", ["OK Energy"], assign=False)
        d = requests.get(f"{API}/segnalazioni/{ok_seg}", headers=admin).json()
        assert d["competenze"] == {"OK Energy": "Andrea"}
        _esito(bo, ok_seg, "KO", assign=False)
        d = requests.get(f"{API}/segnalazioni/{ok_seg}", headers=admin).json()
        assert d["competenze"] == {}
        # back to OK for the export test
        _esito(bo, ok_seg, "OK", ["OK Gas"], assign=False)
        requests.put(f"{API}/segnalazioni/{ok_seg}/competenze", headers=admin, json={"competenze": {"OK Gas": "Rea"}})

    def test_export_has_competenze_column(self, admin, ok_seg):
        r = requests.get(f"{API}/segnalazioni/export/excel", headers=admin, params={"search": "TEST_OkComp"})
        assert r.status_code == 200
        ws = load_workbook(io.BytesIO(r.content)).active
        headers = [c.value for c in ws[1]]
        assert "Competenze" in headers
        idx = headers.index("Competenze")
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
        assert any(r[idx] == "OK Gas: Rea" for r in rows), rows


class TestVisualizzatore:
    def test_sees_all(self, admin, vis):
        a = requests.get(f"{API}/segnalazioni", headers=admin).json()
        v = requests.get(f"{API}/segnalazioni", headers=vis)
        assert v.status_code == 200
        assert {s["id"] for s in a} == {s["id"] for s in v.json()}

    def test_detail_stats_options_segnalatori(self, vis, ok_seg):
        d = requests.get(f"{API}/segnalazioni/{ok_seg}", headers=vis)
        assert d.status_code == 200 and d.json()["competenze"] == {"OK Gas": "Rea"}
        s = requests.get(f"{API}/segnalazioni/stats/summary", headers=vis)
        assert s.status_code == 200 and s.json()["totale"] >= 1
        o = requests.get(f"{API}/segnalazioni/options", headers=vis)
        assert o.status_code == 200 and len(o.json()["commesse"]) >= 1
        sl = requests.get(f"{API}/segnalazioni/segnalatori/list", headers=vis)
        assert sl.status_code == 200 and any(u["username"] == "segnalatore_test" for u in sl.json())
        assert requests.get(f"{API}/segnalazioni/da-lavorare/count", headers=vis).json()["count"] == 0

    def test_write_ops_forbidden(self, vis, ok_seg):
        assert requests.post(f"{API}/segnalazioni", headers=vis, json={"nome": "X", "cognome": "TEST_COMP", "telefono": "3331112222"}).status_code == 403
        assert requests.put(f"{API}/segnalazioni/{ok_seg}/esito", headers=vis, json={"esito": "KO"}).status_code == 403
        assert requests.put(f"{API}/segnalazioni/{ok_seg}/assegnazione", headers=vis, json=ASSIGN).status_code == 403
        assert requests.post(f"{API}/segnalazioni/{ok_seg}/presa-in-carico", headers=vis).status_code == 403
        assert requests.post(f"{API}/segnalazioni/{ok_seg}/documents", headers=vis, files={"file": ("x.txt", b"hi", "text/plain")}).status_code == 403
        assert requests.get(f"{API}/segnalazioni/export/excel", headers=vis).status_code == 403

    def test_download_existing_doc(self, vis, admin):
        docs = [s for s in requests.get(f"{API}/segnalazioni", headers=admin).json() if s.get("documents")]
        if not docs:
            pytest.skip("no docs")
        s = docs[0]
        r_admin = requests.get(f"{API}/segnalazioni/{s['id']}/documents/{s['documents'][0]['id']}/download", headers=admin)
        r_vis = requests.get(f"{API}/segnalazioni/{s['id']}/documents/{s['documents'][0]['id']}/download", headers=vis)
        assert r_vis.status_code == r_admin.status_code, (r_vis.status_code, r_admin.status_code, r_vis.text[:200])
        assert r_vis.status_code == 200

    def test_other_crm_endpoints_not_leaking(self, vis):
        out = {}
        for ep in ["clienti", "users", "commesse", "leads", "sub-agenzie"]:
            r = requests.get(f"{API}/{ep}", headers=vis)
            n = len(r.json()) if r.status_code == 200 and isinstance(r.json(), list) else None
            out[ep] = (r.status_code, n)
        print("VIS other endpoints:", out)


class TestUserCreation:
    def test_admin_creates_visualizzatore(self, admin):
        name = f"test_vis_{uuid.uuid4().hex[:6]}"
        r = requests.post(f"{API}/users", headers=admin, json={"username": name, "email": f"{name}@example.com", "password": "Test12345!", "role": "visualizzatore"})
        assert r.status_code == 200, r.text
        u = r.json()
        assert u["role"] == "visualizzatore"
        r2 = requests.put(f"{API}/users/{u['id']}", headers=admin, json={"username": name, "email": f"{name}@example.com", "role": "visualizzatore"})
        print("update:", r2.status_code, r2.text[:200])
        assert requests.delete(f"{API}/users/{u['id']}", headers=admin).status_code in (200, 204)
